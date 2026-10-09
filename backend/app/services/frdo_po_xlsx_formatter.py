from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.frdo_po_portal_mapping import (
    FRDO_PO_COLUMNS,
    FrdoPoPortalMappingError,
    normalize_frdo_po_classifier_value,
)


FRDO_PO_TEMPLATE_SHEET = (
    "\u0428\u0430\u0431\u043b\u043e\u043d"
)
FRDO_PO_CHECKS_SHEET = (
    "\u041f\u0440\u043e\u0432\u0435\u0440\u043a\u0438"
)
FRDO_PO_OUTPUT_DATE_FORMAT = "dd.mm.yyyy"

FRDO_PO_DOCUMENT_TYPE_TRAINING_REFERENCE = (
    "\u0421\u043f\u0440\u0430\u0432\u043a\u0430 "
    "\u043e\u0431 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u0438"
)
FRDO_PO_DOCUMENT_TYPE_ART_CERTIFICATE = (
    "\u0421\u0432\u0438\u0434\u0435\u0442\u0435\u043b\u044c\u0441\u0442\u0432\u043e "
    "\u043e\u0431 \u043e\u0441\u0432\u043e\u0435\u043d\u0438\u0438 "
    "\u0434\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u044b\u0445 "
    "\u043f\u0440\u0435\u0434\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u044b\u0445 "
    "\u043f\u0440\u043e\u0433\u0440\u0430\u043c\u043c "
    "\u0432 \u043e\u0431\u043b\u0430\u0441\u0442\u0438 "
    "\u0438\u0441\u043a\u0443\u0441\u0441\u0442\u0432"
)
FRDO_PO_DOCUMENT_STATUS_DUPLICATE = (
    "\u0414\u0443\u0431\u043b\u0438\u043a\u0430\u0442"
)
FRDO_PO_NO_VALUE_MARKER = "\u041d\u0435\u0442"

FRDO_PO_GAP_FUNDING_SOURCE_BUDGET = (
    "funding_source_budget"
)
FRDO_PO_GAP_TRAINING_REFERENCE_ART_PREPROFESSIONAL_PROFESSION = (
    "training_reference_art_preprofessional_profession"
)


class FrdoPoXlsxFormatterError(ValueError):
    pass


def _section(
    snapshot: Mapping[str, Any],
    name: str,
) -> Mapping[str, Any]:
    value = snapshot.get(name)

    if not isinstance(value, Mapping):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter snapshot section is invalid: "
            + name
        )

    return value


def _optional_text(
    value: object,
    field: str,
) -> str | None:
    if value is None:
        return None

    if not isinstance(value, str):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter field must be text: "
            + field
        )

    text = value.strip()

    if not text:
        return None

    return text


def _required_text(
    value: object,
    field: str,
) -> str:
    text = _optional_text(
        value,
        field,
    )

    if text is None:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter field is required: "
            + field
        )

    return text


def _classifier(
    classifier: str,
    value: object,
    field: str,
    *,
    required: bool = True,
) -> str | None:
    text = _optional_text(
        value,
        field,
    )

    if text is None:
        if required:
            raise FrdoPoXlsxFormatterError(
                "FRDO PO formatter classifier is required: "
                + field
            )

        return None

    try:
        return normalize_frdo_po_classifier_value(
            classifier,
            text,
        )
    except FrdoPoPortalMappingError as exc:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter classifier is invalid: "
            + field
            + ": "
            + str(exc)
        ) from exc


def _date_value(
    value: object,
    field: str,
) -> date:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if not isinstance(value, str):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter date is required: "
            + field
        )

    raw = value.strip()

    if not raw:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter date is required: "
            + field
        )

    try:
        return date.fromisoformat(
            raw
        )
    except ValueError:
        pass

    try:
        return datetime.fromisoformat(
            raw.replace(
                "Z",
                "+00:00",
            )
        ).date()
    except ValueError:
        pass

    try:
        return datetime.strptime(
            raw,
            "%d.%m.%Y",
        ).date()
    except ValueError as exc:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter date is invalid: "
            + field
        ) from exc


def _hours_value(
    value: object,
) -> int | float:
    if (
        isinstance(value, bool)
        or not isinstance(
            value,
            (
                int,
                float,
            ),
        )
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter course.hours is invalid"
        )

    hours = value

    if hours < 6:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter course.hours must be at least 6"
        )

    if isinstance(
        hours,
        float,
    ) and hours.is_integer():
        return int(hours)

    return hours


def _assert_template_structure(
    workbook: object,
) -> None:
    expected_sheets = [
        FRDO_PO_TEMPLATE_SHEET,
        FRDO_PO_CHECKS_SHEET,
    ]

    if workbook.sheetnames != expected_sheets:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template sheet structure is invalid"
        )

    worksheet = workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]
    checks = workbook[
        FRDO_PO_CHECKS_SHEET
    ]

    if worksheet.max_column != 35:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template column count is invalid"
        )

    if worksheet.max_row != 1002:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template row count is invalid"
        )

    if checks.sheet_state != "hidden":
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template checks sheet must be hidden"
        )

    if (
        len(
            worksheet
            .data_validations
            .dataValidation
        )
        != 20
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template validation structure is invalid"
        )

    expected_headers = [
        item.header
        for item in FRDO_PO_COLUMNS
    ]

    actual_headers = [
        worksheet.cell(
            row=1,
            column=index,
        ).value
        for index in range(
            1,
            36,
        )
    ]

    if actual_headers != expected_headers:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template headers are invalid"
        )

    if any(
        worksheet.cell(
            row=2,
            column=index,
        ).value
        is not None
        for index in range(
            1,
            36,
        )
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template first data row is not empty"
        )


def build_frdo_po_row(
    approval_snapshot: Mapping[str, Any],
) -> tuple[object | None, ...]:
    enrollment = _section(
        approval_snapshot,
        "enrollment",
    )
    course = _section(
        approval_snapshot,
        "course",
    )
    learner = _section(
        approval_snapshot,
        "learner_profile",
    )
    document = _section(
        approval_snapshot,
        "document",
    )
    context = _section(
        approval_snapshot,
        "frdo_context",
    )

    document_type = _classifier(
        "document_type",
        context.get(
            "po_document_type"
        ),
        "frdo_context.po_document_type",
    )
    document_status = _classifier(
        "document_status",
        context.get(
            "document_status"
        ),
        "frdo_context.document_status",
    )
    loss_confirmation = _classifier(
        "loss_confirmation",
        context.get(
            "loss_confirmation"
        ),
        "frdo_context.loss_confirmation",
    )
    exchange_confirmation = _classifier(
        "exchange_confirmation",
        context.get(
            "exchange_confirmation"
        ),
        "frdo_context.exchange_confirmation",
    )
    destruction_confirmation = _classifier(
        "destruction_confirmation",
        context.get(
            "destruction_confirmation"
        ),
        "frdo_context.destruction_confirmation",
    )

    document_series = _optional_text(
        document.get(
            "document_series"
        ),
        "document.document_series",
    )

    if document_series is None:
        if (
            document_type
            != FRDO_PO_DOCUMENT_TYPE_TRAINING_REFERENCE
        ):
            document_series = (
                FRDO_PO_NO_VALUE_MARKER
            )

    document_number = _optional_text(
        document.get(
            "document_number"
        ),
        "document.document_number",
    )

    if (
        document_number is None
        and document_type
        != FRDO_PO_DOCUMENT_TYPE_TRAINING_REFERENCE
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter field is required: "
            "document.document_number"
        )

    issue_date = _date_value(
        document.get(
            "issued_at"
        ),
        "document.issued_at",
    )

    registration_number = _required_text(
        document.get(
            "registration_number"
        ),
        "document.registration_number",
    )

    raw_program_type = _optional_text(
        context.get(
            "po_program_type"
        ),
        "frdo_context.po_program_type",
    )
    raw_profession = _optional_text(
        context.get(
            "po_profession"
        ),
        "frdo_context.po_profession",
    )

    if (
        document_type
        == FRDO_PO_DOCUMENT_TYPE_TRAINING_REFERENCE
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO known mapping gap: "
            + FRDO_PO_GAP_TRAINING_REFERENCE_ART_PREPROFESSIONAL_PROFESSION
        )

    if (
        document_type
        == FRDO_PO_DOCUMENT_TYPE_ART_CERTIFICATE
    ):
        if (
            raw_program_type is not None
            or raw_profession is not None
        ):
            raise FrdoPoXlsxFormatterError(
                "FRDO PO art certificate must have "
                "blank program type and profession"
            )

        program_type = None
        profession = None
    else:
        program_type = _classifier(
            "po_program_type",
            raw_program_type,
            "frdo_context.po_program_type",
            required=False,
        )
        profession = _classifier(
            "po_profession",
            raw_profession,
            "frdo_context.po_profession",
        )

    qualification = _classifier(
        "po_qualification",
        context.get(
            "po_qualification"
        ),
        "frdo_context.po_qualification",
        required=False,
    )

    course_title = _required_text(
        course.get(
            "title"
        ),
        "course.title",
    )

    started_at = _date_value(
        enrollment.get(
            "started_at"
        ),
        "enrollment.started_at",
    )
    completed_at = _date_value(
        enrollment.get(
            "completed_at"
        ),
        "enrollment.completed_at",
    )

    hours = _hours_value(
        course.get(
            "hours"
        )
    )

    last_name = _required_text(
        learner.get(
            "last_name"
        ),
        "learner_profile.last_name",
    )
    first_name = _required_text(
        learner.get(
            "first_name"
        ),
        "learner_profile.first_name",
    )
    middle_name = _optional_text(
        learner.get(
            "middle_name"
        ),
        "learner_profile.middle_name",
    )

    if middle_name is None:
        middle_name = (
            FRDO_PO_NO_VALUE_MARKER
        )

    birth_date = _date_value(
        learner.get(
            "birth_date"
        ),
        "learner_profile.birth_date",
    )

    sex = _classifier(
        "sex",
        learner.get(
            "sex"
        ),
        "learner_profile.sex",
    )

    snils = _optional_text(
        learner.get(
            "snils"
        ),
        "learner_profile.snils",
    )

    citizenship = _classifier(
        "citizenship_country_code",
        learner.get(
            "citizenship_country_code"
        ),
        "learner_profile.citizenship_country_code",
        required=False,
    )

    if (
        issue_date.year >= 2021
        and citizenship is None
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO citizenship is required for "
            "documents issued in 2021 or later"
        )

    if (
        issue_date.year >= 2021
        and citizenship == "643"
        and snils is None
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO SNILS is required for Russian "
            "citizens for documents issued in 2021 or later"
        )

    study_form = _classifier(
        "study_form",
        context.get(
            "study_form"
        ),
        "frdo_context.study_form",
        required=issue_date.year >= 2021,
    )
    raw_funding_source = context.get(
        "funding_source"
    )

    if (
        isinstance(
            raw_funding_source,
            str,
        )
        and raw_funding_source.strip().casefold()
        == "budget"
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO known mapping gap: "
            + FRDO_PO_GAP_FUNDING_SOURCE_BUDGET
        )

    funding_source = _classifier(
        "funding_source",
        raw_funding_source,
        "frdo_context.funding_source",
        required=issue_date.year >= 2021,
    )
    delivery_form = _classifier(
        "education_delivery_form",
        context.get(
            "education_delivery_form"
        ),
        "frdo_context.education_delivery_form",
        required=issue_date.year >= 2021,
    )

    original_values: list[
        object | None
    ] = [
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]

    if (
        document_status
        == FRDO_PO_DOCUMENT_STATUS_DUPLICATE
    ):
        original = context.get(
            "original_document_snapshot_json"
        )

        if not isinstance(
            original,
            Mapping,
        ):
            raise FrdoPoXlsxFormatterError(
                "FRDO PO duplicate original document "
                "snapshot is required"
            )

        original_middle_name = _optional_text(
            original.get(
                "recipient_middle_name"
            ),
            (
                "frdo_context."
                "original_document_snapshot_json."
                "recipient_middle_name"
            ),
        )

        if original_middle_name is None:
            original_middle_name = (
                FRDO_PO_NO_VALUE_MARKER
            )

        original_values = [
            _classifier(
                "document_type",
                original.get(
                    "document_type"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "document_type"
                ),
            ),
            _required_text(
                original.get(
                    "document_series"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "document_series"
                ),
            ),
            _required_text(
                original.get(
                    "document_number"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "document_number"
                ),
            ),
            _required_text(
                original.get(
                    "registration_number"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "registration_number"
                ),
            ),
            _date_value(
                original.get(
                    "issue_date"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "issue_date"
                ),
            ),
            _required_text(
                original.get(
                    "recipient_last_name"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "recipient_last_name"
                ),
            ),
            _required_text(
                original.get(
                    "recipient_first_name"
                ),
                (
                    "frdo_context."
                    "original_document_snapshot_json."
                    "recipient_first_name"
                ),
            ),
            original_middle_name,
        ]

    row: tuple[
        object | None,
        ...,
    ] = (
        document_type,
        document_status,
        loss_confirmation,
        exchange_confirmation,
        destruction_confirmation,
        document_series,
        document_number,
        issue_date,
        registration_number,
        program_type,
        course_title,
        profession,
        qualification,
        started_at.year,
        completed_at.year,
        hours,
        last_name,
        first_name,
        middle_name,
        birth_date,
        sex,
        snils,
        citizenship,
        study_form,
        funding_source,
        delivery_form,
        *original_values,
        None,
    )

    if len(row) != len(
        FRDO_PO_COLUMNS
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter row width is invalid"
        )

    return row


FRDO_PO_BATCH_MAX_RECORDS = 1001


def format_frdo_po_batch_xlsx(
    *,
    approval_snapshots: Sequence[Mapping[str, Any]],
    template_bytes: bytes,
) -> bytes:
    if (
        not isinstance(approval_snapshots, Sequence)
        or isinstance(
            approval_snapshots,
            (str, bytes, bytearray),
        )
    ):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO batch snapshots must be a sequence"
        )

    record_count = len(approval_snapshots)

    if not 1 <= record_count <= FRDO_PO_BATCH_MAX_RECORDS:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO batch record count must be between 1 and "
            + str(FRDO_PO_BATCH_MAX_RECORDS)
        )

    if not isinstance(template_bytes, bytes) or not template_bytes:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter template bytes are invalid"
        )

    rows: list[tuple[object | None, ...]] = []

    for position, snapshot in enumerate(approval_snapshots):
        if not isinstance(snapshot, Mapping):
            raise FrdoPoXlsxFormatterError(
                "FRDO PO batch snapshot must be a mapping at position "
                + str(position)
            )

        rows.append(build_frdo_po_row(snapshot))

    try:
        workbook = load_workbook(
            BytesIO(template_bytes),
            read_only=False,
            data_only=False,
        )
    except Exception as exc:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO template cannot be opened"
        ) from exc

    _assert_template_structure(workbook)

    worksheet = workbook[FRDO_PO_TEMPLATE_SHEET]

    for position in range(record_count):
        row_number = position + 2

        if any(
            worksheet.cell(
                row=row_number,
                column=column,
            ).value is not None
            for column in range(1, 36)
        ):
            raise FrdoPoXlsxFormatterError(
                "FRDO PO template data row is not empty: "
                + str(row_number)
            )

    for position, row_values in enumerate(rows):
        row_number = position + 2

        for column, value in enumerate(row_values, start=1):
            worksheet.cell(
                row=row_number,
                column=column,
            ).value = value

        for column in (8, 20, 31):
            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            if cell.value is not None:
                cell.number_format = FRDO_PO_OUTPUT_DATE_FORMAT

    buffer = BytesIO()

    try:
        workbook.save(buffer)
    except Exception as exc:
        raise FrdoPoXlsxFormatterError(
            "FRDO PO XLSX serialization failed"
        ) from exc

    content = buffer.getvalue()

    if not content.startswith(b"PK"):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO XLSX serialization produced an invalid package"
        )

    return content


def format_frdo_po_xlsx(
    *,
    approval_snapshot: Mapping[str, Any],
    template_bytes: bytes,
) -> bytes:
    if not isinstance(approval_snapshot, Mapping):
        raise FrdoPoXlsxFormatterError(
            "FRDO PO formatter snapshot must be a mapping"
        )

    return format_frdo_po_batch_xlsx(
        approval_snapshots=(approval_snapshot,),
        template_bytes=template_bytes,
    )
