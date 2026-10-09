from __future__ import annotations

from collections.abc import Mapping, Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.frdo_po_xlsx_formatter import (
    FRDO_PO_CHECKS_SHEET,
    FRDO_PO_OUTPUT_DATE_FORMAT,
    FRDO_PO_TEMPLATE_SHEET,
    FrdoPoXlsxFormatterError,
    FRDO_PO_BATCH_MAX_RECORDS,
    build_frdo_po_row,
)


class FrdoPoXlsxValidatorError(ValueError):
    pass


def _load_xlsx(
    content: object,
    *,
    label: str,
):
    if (
        not isinstance(
            content,
            bytes,
        )
        or not content
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO "
            + label
            + " must be non-empty XLSX bytes"
        )

    if not content.startswith(
        b"PK"
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO "
            + label
            + " is not an XLSX package"
        )

    try:
        return load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
        )
    except Exception as exc:
        raise FrdoPoXlsxValidatorError(
            "FRDO PO "
            + label
            + " cannot be opened"
        ) from exc


def _defined_name_signature(
    workbook: object,
) -> tuple[
    tuple[
        object,
        ...,
    ],
    ...,
]:
    return tuple(
        (
            item.name,
            item.attr_text,
            item.hidden,
            item.localSheetId,
        )
        for item
        in workbook.defined_names.values()
    )


def _validation_signature(
    worksheet: object,
) -> tuple[
    tuple[
        object,
        ...,
    ],
    ...,
]:
    return tuple(
        (
            item.type,
            item.formula1,
            item.formula2,
            str(item.sqref),
            item.allowBlank,
            item.operator,
            item.showErrorMessage,
            item.errorTitle,
            item.error,
        )
        for item
        in (
            worksheet
            .data_validations
            .dataValidation
        )
    )


def _style_signature(
    cell: object,
) -> tuple[
    object,
    ...,
]:
    style = cell._style

    return (
        style.fontId,
        style.fillId,
        style.borderId,
        style.alignmentId,
        style.protectionId,
        style.pivotButton,
        style.quotePrefix,
    )


def _date_only(
    value: object,
) -> date | None:
    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if isinstance(
        value,
        date,
    ):
        return value

    return None


def _compare_sheet_values(
    *,
    actual: object,
    expected: object,
    start_row: int,
    end_row: int,
    start_column: int,
    end_column: int,
    label: str,
) -> None:
    for row in range(
        start_row,
        end_row + 1,
    ):
        for column in range(
            start_column,
            end_column + 1,
        ):
            actual_value = actual.cell(
                row=row,
                column=column,
            ).value

            expected_value = expected.cell(
                row=row,
                column=column,
            ).value

            if actual_value != expected_value:
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact changed "
                    + label
                    + " cell "
                    + actual.cell(
                        row=row,
                        column=column,
                    ).coordinate
                )


def validate_frdo_po_batch_xlsx(
    *,
    content: bytes,
    approval_snapshots: Sequence[Mapping[str, Any]],
    template_bytes: bytes,
) -> None:
    if (
        not isinstance(approval_snapshots, Sequence)
        or isinstance(
            approval_snapshots,
            (str, bytes, bytearray),
        )
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO batch snapshots must be a sequence"
        )

    count = len(approval_snapshots)

    if not 1 <= count <= FRDO_PO_BATCH_MAX_RECORDS:
        raise FrdoPoXlsxValidatorError(
            "FRDO PO batch record count must be between 1 and "
            + str(FRDO_PO_BATCH_MAX_RECORDS)
        )

    expected_rows: list[tuple[object | None, ...]] = []

    for position, snapshot in enumerate(approval_snapshots):
        if not isinstance(snapshot, Mapping):
            raise FrdoPoXlsxValidatorError(
                "FRDO PO batch snapshot must be a mapping at position "
                + str(position)
            )

        try:
            expected_row = build_frdo_po_row(snapshot)
        except FrdoPoXlsxFormatterError as exc:
            raise FrdoPoXlsxValidatorError(
                "FRDO PO validator snapshot cannot build expected row: "
                + str(exc)
            ) from exc

        if len(expected_row) != 35:
            raise FrdoPoXlsxValidatorError(
                "FRDO PO expected row width is invalid"
            )

        expected_rows.append(expected_row)

    artifact_workbook = _load_xlsx(
        content,
        label="artifact",
    )
    template_workbook = _load_xlsx(
        template_bytes,
        label="template",
    )

    if (
        artifact_workbook.sheetnames
        != template_workbook.sheetnames
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact sheet list changed"
        )

    if artifact_workbook.sheetnames != [
        FRDO_PO_TEMPLATE_SHEET,
        FRDO_PO_CHECKS_SHEET,
    ]:
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact sheet structure is unsupported"
        )

    for sheet_name in artifact_workbook.sheetnames:
        artifact_sheet = artifact_workbook[
            sheet_name
        ]
        template_sheet = template_workbook[
            sheet_name
        ]

        if (
            artifact_sheet.sheet_state
            != template_sheet.sheet_state
        ):
            raise FrdoPoXlsxValidatorError(
                "FRDO PO artifact sheet state changed: "
                + sheet_name
            )

    artifact_sheet = artifact_workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]
    template_sheet = template_workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]

    artifact_checks = artifact_workbook[
        FRDO_PO_CHECKS_SHEET
    ]
    template_checks = template_workbook[
        FRDO_PO_CHECKS_SHEET
    ]

    if (
        artifact_sheet.max_row
        != template_sheet.max_row
        or artifact_sheet.max_column
        != template_sheet.max_column
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact template dimensions changed"
        )

    if (
        artifact_checks.max_row
        != template_checks.max_row
        or artifact_checks.max_column
        != template_checks.max_column
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact checks dimensions changed"
        )

    if (
        tuple(
            str(item)
            for item
            in artifact_sheet.merged_cells.ranges
        )
        != tuple(
            str(item)
            for item
            in template_sheet.merged_cells.ranges
        )
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact merged cells changed"
        )

    if (
        artifact_sheet.freeze_panes
        != template_sheet.freeze_panes
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact freeze panes changed"
        )

    if (
        artifact_sheet.auto_filter.ref
        != template_sheet.auto_filter.ref
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact auto filter changed"
        )

    if (
        _validation_signature(
            artifact_sheet
        )
        != _validation_signature(
            template_sheet
        )
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact data validation metadata changed"
        )

    if (
        _defined_name_signature(
            artifact_workbook
        )
        != _defined_name_signature(
            template_workbook
        )
    ):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO artifact defined names changed"
        )

    for column in range(
        1,
        36,
    ):
        artifact_header = artifact_sheet.cell(
            row=1,
            column=column,
        ).value

        template_header = template_sheet.cell(
            row=1,
            column=column,
        ).value

        if artifact_header != template_header:
            raise FrdoPoXlsxValidatorError(
                "FRDO PO artifact header changed at "
                + artifact_sheet.cell(
                    row=1,
                    column=column,
                ).coordinate
            )

    for position, expected_row in enumerate(expected_rows):
        row_number = position + 2

        for column, expected_value in enumerate(
            expected_row,
            start=1,
        ):
            cell = artifact_sheet.cell(
                row=row_number,
                column=column,
            )

            actual_value = cell.value

            if isinstance(expected_value, date):
                if cell.number_format != FRDO_PO_OUTPUT_DATE_FORMAT:
                    raise FrdoPoXlsxValidatorError(
                        "FRDO PO artifact date format mismatch at "
                        + cell.coordinate
                    )

                actual_date = _date_only(actual_value)

                if actual_date != expected_value:
                    raise FrdoPoXlsxValidatorError(
                        "FRDO PO artifact value mismatch at "
                        + cell.coordinate
                    )

            elif actual_value != expected_value:
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact value mismatch at "
                    + cell.coordinate
                )

            template_cell = template_sheet.cell(
                row=row_number,
                column=column,
            )

            if _style_signature(cell) != _style_signature(
                template_cell
            ):
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact style changed at "
                    + cell.coordinate
                )

            if (
                not isinstance(expected_value, date)
                and cell.number_format
                != template_cell.number_format
            ):
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact number format changed at "
                    + cell.coordinate
                )

    _compare_sheet_values(
        actual=artifact_sheet,
        expected=template_sheet,
        start_row=2 + len(expected_rows),
        end_row=template_sheet.max_row,
        start_column=1,
        end_column=template_sheet.max_column,
        label="template",
    )

    # Check that unused cells and header retain their template styles.
    main_sheet_max_row = template_sheet.max_row
    main_sheet_max_column = template_sheet.max_column
    for row_number in range(1, main_sheet_max_row + 1):
        if 2 <= row_number < 2 + len(expected_rows):
            continue

        for column in range(1, main_sheet_max_column + 1):
            actual_cell = artifact_sheet.cell(
                row=row_number,
                column=column,
            )
            template_cell = template_sheet.cell(
                row=row_number,
                column=column,
            )

            actual_style = (
                _style_signature(actual_cell)
                if actual_cell._style is not None
                else None
            )
            expected_style = (
                _style_signature(template_cell)
                if template_cell._style is not None
                else None
            )

            if (
                actual_style != expected_style
                or actual_cell.number_format
                != template_cell.number_format
            ):
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact style changed at "
                    + actual_cell.coordinate
                )

    _compare_sheet_values(
        actual=artifact_checks,
        expected=template_checks,
        start_row=1,
        end_row=template_checks.max_row,
        start_column=1,
        end_column=template_checks.max_column,
        label="checks",
    )

    # The hidden checks sheet must retain its formatting.
    checks_sheet_max_row = template_checks.max_row
    checks_sheet_max_column = template_checks.max_column
    for row_number in range(1, checks_sheet_max_row + 1):
        for column in range(1, checks_sheet_max_column + 1):
            actual_cell = artifact_checks.cell(
                row=row_number,
                column=column,
            )
            template_cell = template_checks.cell(
                row=row_number,
                column=column,
            )

            actual_style = (
                _style_signature(actual_cell)
                if actual_cell._style is not None
                else None
            )
            expected_style = (
                _style_signature(template_cell)
                if template_cell._style is not None
                else None
            )

            if (
                actual_style != expected_style
                or actual_cell.number_format
                != template_cell.number_format
            ):
                raise FrdoPoXlsxValidatorError(
                    "FRDO PO artifact checks style changed at "
                    + actual_cell.coordinate
                )



def validate_frdo_po_xlsx(
    *,
    content: bytes,
    approval_snapshot: Mapping[str, Any],
    template_bytes: bytes,
) -> None:
    if not isinstance(approval_snapshot, Mapping):
        raise FrdoPoXlsxValidatorError(
            "FRDO PO validator snapshot must be a mapping"
        )

    validate_frdo_po_batch_xlsx(
        content=content,
        approval_snapshots=(approval_snapshot,),
        template_bytes=template_bytes,
    )
