from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.compliance_registry_approval import (
    FRDO_PO_APPROVAL_SNAPSHOT_SCHEMA_VERSION,
)
from app.services.compliance_registry_contract import (
    PROGRAM_TYPE_VOCATIONAL_TRAINING,
)
from app.services.frdo_po_portal_contract import (
    read_verified_frdo_po_template_bytes,
)
from app.services.frdo_po_xlsx_formatter import (
    FRDO_PO_CHECKS_SHEET,
    FRDO_PO_OUTPUT_DATE_FORMAT,
    FRDO_PO_TEMPLATE_SHEET,
    FrdoPoXlsxFormatterError,
    build_frdo_po_row,
    format_frdo_po_xlsx,
    format_frdo_po_batch_xlsx,
)


DOCUMENT_TYPE = (
    "\u0421\u0432\u0438\u0434\u0435\u0442\u0435\u043b\u044c\u0441\u0442\u0432\u043e "
    "\u043e \u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u0438 "
    "\u0440\u0430\u0431\u043e\u0447\u0435\u0433\u043e, "
    "\u0434\u043e\u043b\u0436\u043d\u043e\u0441\u0442\u0438 "
    "\u0441\u043b\u0443\u0436\u0430\u0449\u0435\u0433\u043e"
)
TRAINING_REFERENCE = (
    "\u0421\u043f\u0440\u0430\u0432\u043a\u0430 "
    "\u043e\u0431 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u0438"
)
ART_CERTIFICATE = (
    "\u0421\u0432\u0438\u0434\u0435\u0442\u0435\u043b\u044c\u0441\u0442\u0432\u043e "
    "\u043e\u0431 \u043e\u0441\u0432\u043e\u0435\u043d\u0438\u0438 "
    "\u0434\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u044b\u0445 "
    "\u043f\u0440\u0435\u0434\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u044b\u0445 "
    "\u043f\u0440\u043e\u0433\u0440\u0430\u043c\u043c "
    "\u0432 \u043e\u0431\u043b\u0430\u0441\u0442\u0438 "
    "\u0438\u0441\u043a\u0443\u0441\u0441\u0442\u0432"
)
PROGRAM_TYPE = (
    "\u041f\u0440\u043e\u0433\u0440\u0430\u043c\u043c\u0430 "
    "\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u043e\u0439 "
    "\u043f\u043e\u0434\u0433\u043e\u0442\u043e\u0432\u043a\u0438 "
    "\u043f\u043e \u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u0438 "
    "\u0440\u0430\u0431\u043e\u0447\u0435\u0433\u043e, "
    "\u0434\u043e\u043b\u0436\u043d\u043e\u0441\u0442\u0438 "
    "\u0441\u043b\u0443\u0436\u0430\u0449\u0435\u0433\u043e"
)
PROFESSION = (
    "\u0412\u043e\u0434\u0438\u0442\u0435\u043b\u044c "
    "\u0430\u0432\u0442\u043e\u043c\u043e\u0431\u0438\u043b\u044f"
)


def make_snapshot() -> dict:
    return {
        "schema_version": (
            FRDO_PO_APPROVAL_SNAPSHOT_SCHEMA_VERSION
        ),
        "registry": "frdo",
        "enrollment": {
            "status": "completed",
            "started_at": "2026-09-01T00:00:00+00:00",
            "completed_at": "2026-09-30T00:00:00+00:00",
        },
        "course": {
            "title": (
                "\u041e\u0431\u0443\u0447\u0435\u043d\u0438\u0435 "
                "\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u0435\u0439"
            ),
            "hours": 40,
            "regulatory_program_type": (
                PROGRAM_TYPE_VOCATIONAL_TRAINING
            ),
        },
        "learner_profile": {
            "last_name": "\u0418\u0432\u0430\u043d\u043e\u0432",
            "first_name": "\u0418\u0432\u0430\u043d",
            "middle_name": "\u0418\u0432\u0430\u043d\u043e\u0432\u0438\u0447",
            "birth_date": "2000-01-01",
            "sex": "male",
            "snils": "123-456-789 00",
            "citizenship_country_code": "643",
        },
        "document": {
            "enrollment_id": "enrollment-1",
            "document_series": "PO",
            "document_number": "1",
            "document_type": "certificate",
            "issued_at": "2026-09-30",
            "registration_number": "REG-1",
            "revoked_at": None,
        },
        "frdo_context": {
            "document_status": "original",
            "loss_confirmation": "No",
            "exchange_confirmation": "No",
            "destruction_confirmation": "No",
            "study_form": "full_time",
            "funding_source": "paid",
            "education_delivery_form": "onsite",
            "po_document_type": DOCUMENT_TYPE,
            "po_program_type": "initial_training",
            "po_profession": PROFESSION,
            "po_qualification": "1",
            "original_document_snapshot_json": None,
        },
    }


def template_bytes() -> bytes:
    return read_verified_frdo_po_template_bytes(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )


def test_formatter_fills_exact_first_data_row() -> None:
    content = format_frdo_po_xlsx(
        approval_snapshot=make_snapshot(),
        template_bytes=template_bytes(),
    )

    workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=False,
    )
    sheet = workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]

    assert workbook.sheetnames == [
        FRDO_PO_TEMPLATE_SHEET,
        FRDO_PO_CHECKS_SHEET,
    ]
    assert (
        workbook[
            FRDO_PO_CHECKS_SHEET
        ].sheet_state
        == "hidden"
    )
    assert sheet.max_column == 35
    assert sheet.max_row == 1002
    assert (
        len(
            sheet
            .data_validations
            .dataValidation
        )
        == 20
    )

    assert sheet["A2"].value == DOCUMENT_TYPE
    assert sheet["B2"].value == "\u041e\u0440\u0438\u0433\u0438\u043d\u0430\u043b"
    assert sheet["C2"].value == "\u041d\u0435\u0442"
    assert sheet["D2"].value == "\u041d\u0435\u0442"
    assert sheet["E2"].value == "\u041d\u0435\u0442"
    assert sheet["F2"].value == "PO"
    assert sheet["G2"].value == "1"
    assert sheet["I2"].value == "REG-1"
    assert sheet["J2"].value == PROGRAM_TYPE
    assert (
        sheet["K2"].value
        == "\u041e\u0431\u0443\u0447\u0435\u043d\u0438\u0435 \u0432\u043e\u0434\u0438\u0442\u0435\u043b\u0435\u0439"
    )
    assert sheet["L2"].value == PROFESSION
    assert sheet["M2"].value == "1"
    assert sheet["N2"].value == 2026
    assert sheet["O2"].value == 2026
    assert sheet["P2"].value == 40
    assert sheet["Q2"].value == "\u0418\u0432\u0430\u043d\u043e\u0432"
    assert sheet["R2"].value == "\u0418\u0432\u0430\u043d"
    assert sheet["S2"].value == "\u0418\u0432\u0430\u043d\u043e\u0432\u0438\u0447"
    assert sheet["U2"].value == "\u041c\u0443\u0436"
    assert sheet["V2"].value == "123-456-789 00"
    assert sheet["W2"].value == "643"
    assert sheet["X2"].value == "\u041e\u0447\u043d\u0430\u044f"
    assert (
        sheet["Y2"].value
        == "\u041f\u043b\u0430\u0442\u043d\u043e\u0435 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u0435"
    )
    assert (
        sheet["Z2"].value
        == "\u0432 \u043e\u0431\u0440\u0430\u0437\u043e\u0432\u0430\u0442\u0435\u043b\u044c\u043d\u043e\u0439 \u043e\u0440\u0433\u0430\u043d\u0438\u0437\u0430\u0446\u0438\u0438"
    )

    assert sheet["H2"].value.date().isoformat() == "2026-09-30"
    assert sheet["T2"].value.date().isoformat() == "2000-01-01"
    assert sheet["H2"].number_format == FRDO_PO_OUTPUT_DATE_FORMAT
    assert sheet["T2"].number_format == FRDO_PO_OUTPUT_DATE_FORMAT

    for column in (
        "AA",
        "AB",
        "AC",
        "AD",
        "AE",
        "AF",
        "AG",
        "AH",
        "AI",
    ):
        assert sheet[
            column + "2"
        ].value is None


def test_formatter_populates_duplicate_original_fields() -> None:
    snapshot = make_snapshot()

    snapshot[
        "frdo_context"
    ][
        "document_status"
    ] = "duplicate"

    snapshot[
        "frdo_context"
    ][
        "original_document_snapshot_json"
    ] = {
        "document_type": DOCUMENT_TYPE,
        "document_series": "OLD",
        "document_number": "12345",
        "registration_number": "REG-OLD",
        "issue_date": "2025-05-06",
        "recipient_last_name": "\u041f\u0435\u0442\u0440\u043e\u0432",
        "recipient_first_name": "\u041f\u0435\u0442\u0440",
        "recipient_middle_name": None,
    }

    content = format_frdo_po_xlsx(
        approval_snapshot=snapshot,
        template_bytes=template_bytes(),
    )

    workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=False,
    )
    sheet = workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]

    assert sheet["B2"].value == "\u0414\u0443\u0431\u043b\u0438\u043a\u0430\u0442"
    assert sheet["AA2"].value == DOCUMENT_TYPE
    assert sheet["AB2"].value == "OLD"
    assert sheet["AC2"].value == "12345"
    assert sheet["AD2"].value == "REG-OLD"
    assert sheet["AE2"].value.date().isoformat() == "2025-05-06"
    assert sheet["AE2"].number_format == FRDO_PO_OUTPUT_DATE_FORMAT
    assert sheet["AF2"].value == "\u041f\u0435\u0442\u0440\u043e\u0432"
    assert sheet["AG2"].value == "\u041f\u0435\u0442\u0440"
    assert sheet["AH2"].value == "\u041d\u0435\u0442"
    assert sheet["AI2"].value is None


def test_formatter_uses_no_patronymic_marker() -> None:
    snapshot = make_snapshot()
    snapshot[
        "learner_profile"
    ][
        "middle_name"
    ] = None

    row = build_frdo_po_row(
        snapshot
    )

    assert row[18] == "\u041d\u0435\u0442"


def test_training_reference_known_gap_fails_closed() -> None:
    snapshot = make_snapshot()
    snapshot[
        "frdo_context"
    ][
        "po_document_type"
    ] = TRAINING_REFERENCE
    snapshot[
        "document"
    ][
        "document_series"
    ] = None
    snapshot[
        "document"
    ][
        "document_number"
    ] = None

    with pytest.raises(
        FrdoPoXlsxFormatterError,
        match=(
            "training_reference_"
            "art_preprofessional_profession"
        ),
    ):
        build_frdo_po_row(
            snapshot
        )


def test_missing_series_uses_official_no_value_marker() -> None:
    snapshot = make_snapshot()
    snapshot[
        "document"
    ][
        "document_series"
    ] = None

    row = build_frdo_po_row(
        snapshot
    )

    assert row[5] == "\u041d\u0435\u0442"


def test_generic_budget_fails_closed() -> None:
    snapshot = make_snapshot()
    snapshot[
        "frdo_context"
    ][
        "funding_source"
    ] = "budget"

    with pytest.raises(
        FrdoPoXlsxFormatterError,
        match="funding_source_budget",
    ):
        build_frdo_po_row(
            snapshot
        )


def test_art_certificate_rejects_nonblank_program_fields() -> None:
    snapshot = make_snapshot()
    snapshot[
        "frdo_context"
    ][
        "po_document_type"
    ] = ART_CERTIFICATE

    with pytest.raises(
        FrdoPoXlsxFormatterError,
        match="art certificate",
    ):
        build_frdo_po_row(
            snapshot
        )


def test_template_header_change_fails_closed() -> None:
    workbook = load_workbook(
        BytesIO(
            template_bytes()
        ),
        read_only=False,
        data_only=False,
    )

    workbook[
        FRDO_PO_TEMPLATE_SHEET
    ][
        "A1"
    ] = "changed"

    changed = BytesIO()
    workbook.save(
        changed
    )

    with pytest.raises(
        FrdoPoXlsxFormatterError,
        match="headers",
    ):
        format_frdo_po_xlsx(
            approval_snapshot=make_snapshot(),
            template_bytes=changed.getvalue(),
        )



def test_formatter_preserves_template_structure_and_styles() -> None:
    source_bytes = template_bytes()

    source_workbook = load_workbook(
        BytesIO(source_bytes),
        read_only=False,
        data_only=False,
    )

    source_sheet = source_workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]

    content = format_frdo_po_xlsx(
        approval_snapshot=make_snapshot(),
        template_bytes=source_bytes,
    )

    generated_workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=False,
    )

    generated_sheet = generated_workbook[
        FRDO_PO_TEMPLATE_SHEET
    ]

    assert (
        generated_workbook.sheetnames
        == source_workbook.sheetnames
    )

    assert (
        generated_workbook[
            FRDO_PO_CHECKS_SHEET
        ].sheet_state
        == "hidden"
    )

    assert (
        len(
            generated_sheet
            .data_validations
            .dataValidation
        )
        == len(
            source_sheet
            .data_validations
            .dataValidation
        )
    )

    assert [
        item.name
        for item
        in generated_workbook.defined_names.values()
    ] == [
        item.name
        for item
        in source_workbook.defined_names.values()
    ]

    def style_signature(cell):
        style = cell._style

        return {
            "fontId": style.fontId,
            "fillId": style.fillId,
            "borderId": style.borderId,
            "alignmentId": style.alignmentId,
            "protectionId": style.protectionId,
            "pivotButton": style.pivotButton,
            "quotePrefix": style.quotePrefix,
        }

    for column in range(
        1,
        36,
    ):
        source_cell = source_sheet.cell(
            row=2,
            column=column,
        )
        generated_cell = generated_sheet.cell(
            row=2,
            column=column,
        )

        assert (
            style_signature(
                generated_cell
            )
            == style_signature(
                source_cell
            )
        )

        if column in (
            8,
            20,
        ):
            assert (
                generated_cell.number_format
                == FRDO_PO_OUTPUT_DATE_FORMAT
            )
        else:
            assert (
                generated_cell.number_format
                == source_cell.number_format
            )

    for row in range(
        3,
        1003,
    ):
        for column in range(
            1,
            36,
        ):
            assert (
                generated_sheet.cell(
                    row=row,
                    column=column,
                ).value
                is None
            )

    assert (
        template_bytes()
        == source_bytes
    )

def test_formatter_does_not_mutate_frozen_snapshot() -> None:
    snapshot = make_snapshot()
    before = deepcopy(
        snapshot
    )

    format_frdo_po_xlsx(
        approval_snapshot=snapshot,
        template_bytes=template_bytes(),
    )

    assert snapshot == before


def _batch_test_snapshots(count: int) -> list[dict]:
    snapshots = []

    for index in range(count):
        snapshot = make_snapshot()

        snapshot["document"]["document_number"] = str(index + 1)
        snapshot["document"]["registration_number"] = (
            "REG-" + str(index + 1)
        )

        snapshots.append(snapshot)

    return snapshots


def _assert_batch_test_row(sheet, row_number, snapshot):
    expected = build_frdo_po_row(snapshot)

    assert len(expected) == 35

    for column, value in enumerate(expected, start=1):
        cell = sheet.cell(
            row=row_number,
            column=column,
        )

        actual = cell.value

        if (
            isinstance(actual, datetime)
            and isinstance(value, date)
            and not isinstance(value, datetime)
        ):
            actual = actual.date()

        assert actual == value, (
            "Mismatch at row "
            + str(row_number)
            + ", column "
            + str(column)
        )

        if column in (8, 20, 31) and value is not None:
            assert cell.number_format == FRDO_PO_OUTPUT_DATE_FORMAT


def test_batch_formatter_one_record_matches_single_flow():
    snapshot = make_snapshot()
    source = template_bytes()

    single = format_frdo_po_xlsx(
        approval_snapshot=snapshot,
        template_bytes=source,
    )

    batch = format_frdo_po_batch_xlsx(
        approval_snapshots=[snapshot],
        template_bytes=source,
    )

    single_book = load_workbook(BytesIO(single))
    batch_book = load_workbook(BytesIO(batch))

    single_sheet = single_book[FRDO_PO_TEMPLATE_SHEET]
    batch_sheet = batch_book[FRDO_PO_TEMPLATE_SHEET]

    _assert_batch_test_row(batch_sheet, 2, snapshot)

    for column in range(1, 36):
        left = single_sheet.cell(row=2, column=column)
        right = batch_sheet.cell(row=2, column=column)

        assert left.value == right.value
        assert left.number_format == right.number_format
        assert left.style_id == right.style_id

    assert batch_book.sheetnames == single_book.sheetnames
    assert batch_book[FRDO_PO_CHECKS_SHEET].sheet_state == "hidden"

    assert all(
        batch_sheet.cell(row=3, column=column).value is None
        for column in range(1, 36)
    )


def test_batch_formatter_two_records_preserves_order():
    snapshots = _batch_test_snapshots(2)
    original = deepcopy(snapshots)

    content = format_frdo_po_batch_xlsx(
        approval_snapshots=snapshots,
        template_bytes=template_bytes(),
    )

    workbook = load_workbook(BytesIO(content))
    sheet = workbook[FRDO_PO_TEMPLATE_SHEET]

    _assert_batch_test_row(sheet, 2, snapshots[0])
    _assert_batch_test_row(sheet, 3, snapshots[1])

    assert all(
        sheet.cell(row=4, column=column).value is None
        for column in range(1, 36)
    )

    assert snapshots == original
    assert workbook.sheetnames == [
        FRDO_PO_TEMPLATE_SHEET,
        FRDO_PO_CHECKS_SHEET,
    ]
    assert workbook[FRDO_PO_CHECKS_SHEET].sheet_state == "hidden"


def test_batch_formatter_1001_records():
    snapshots = _batch_test_snapshots(1001)

    content = format_frdo_po_batch_xlsx(
        approval_snapshots=snapshots,
        template_bytes=template_bytes(),
    )

    workbook = load_workbook(BytesIO(content))
    sheet = workbook[FRDO_PO_TEMPLATE_SHEET]

    assert sheet.max_row == 1002
    assert sheet.max_column == 35

    _assert_batch_test_row(sheet, 2, snapshots[0])
    _assert_batch_test_row(sheet, 1002, snapshots[-1])

    assert workbook[FRDO_PO_CHECKS_SHEET].sheet_state == "hidden"
    assert len(sheet.data_validations.dataValidation) == 20


def test_batch_formatter_rejects_1002_records():
    snapshots = [make_snapshot()] * 1002

    with pytest.raises(FrdoPoXlsxFormatterError):
        format_frdo_po_batch_xlsx(
            approval_snapshots=snapshots,
            template_bytes=template_bytes(),
        )


def test_batch_formatter_rejects_empty_sequence():
    with pytest.raises(FrdoPoXlsxFormatterError):
        format_frdo_po_batch_xlsx(
            approval_snapshots=[],
            template_bytes=template_bytes(),
        )


def test_batch_formatter_rejects_invalid_snapshot():
    snapshots = [make_snapshot(), None]

    with pytest.raises(FrdoPoXlsxFormatterError):
        format_frdo_po_batch_xlsx(
            approval_snapshots=snapshots,
            template_bytes=template_bytes(),
        )


def test_batch_formatter_second_record_known_gap_fails_closed():
    snapshots = _batch_test_snapshots(2)

    snapshots[1]["frdo_context"]["funding_source"] = "budget"

    with pytest.raises(FrdoPoXlsxFormatterError):
        format_frdo_po_batch_xlsx(
            approval_snapshots=snapshots,
            template_bytes=template_bytes(),
        )
