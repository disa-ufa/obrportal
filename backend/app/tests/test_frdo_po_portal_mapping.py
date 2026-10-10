from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.services.compliance_registry_contract import (
    PROGRAM_TYPE_VOCATIONAL_TRAINING,
)
from app.services.frdo_po_portal_contract import (
    read_verified_frdo_po_template_bytes,
)
from app.services.frdo_po_portal_mapping import (
    CATEGORY_CONSTANT,
    FRDO_PO_COLUMNS,
    FRDO_PO_KNOWN_MAPPING_GAPS,
    FRDO_PO_MAPPING_VERSION,
    FrdoPoPortalMappingError,
    get_frdo_po_classifier_values,
    normalize_frdo_po_classifier_value,
)


def test_mapping_contract_has_all_35_columns() -> None:
    assert FRDO_PO_MAPPING_VERSION == (
        "frdo-po-working-reference-mapping-v1"
    )
    assert len(FRDO_PO_COLUMNS) == 35
    assert [
        item.index
        for item in FRDO_PO_COLUMNS
    ] == list(range(1, 36))

    assert FRDO_PO_COLUMNS[0].letter == "A"
    assert FRDO_PO_COLUMNS[-1].letter == "AI"


def test_mapping_headers_match_pinned_workbook() -> None:
    content = read_verified_frdo_po_template_bytes(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=False,
    )

    worksheet = workbook[
        "\u0428\u0430\u0431\u043b\u043e\u043d"
    ]

    actual = tuple(
        str(
            worksheet.cell(
                row=1,
                column=item.index,
            ).value
        )
        for item in FRDO_PO_COLUMNS
    )

    expected = tuple(
        item.header
        for item in FRDO_PO_COLUMNS
    )

    assert actual == expected


@pytest.mark.parametrize(
    ("classifier", "expected_count"),
    [
        ("document_type", 3),
        ("document_status", 2),
        ("loss_confirmation", 5),
        ("exchange_confirmation", 3),
        ("destruction_confirmation", 2),
        ("po_program_type", 3),
        ("po_profession", 5599),
        ("po_qualification", 34),
        ("sex", 2),
        ("citizenship_country_code", 254),
        ("study_form", 3),
        ("funding_source", 4),
        ("education_delivery_form", 2),
    ],
)
def test_pinned_classifier_counts(
    classifier: str,
    expected_count: int,
) -> None:
    values = get_frdo_po_classifier_values(
        classifier
    )

    assert len(values) == expected_count


@pytest.mark.parametrize(
    ("classifier", "source", "expected"),
    [
        (
            "document_status",
            "Original",
            "\u041e\u0440\u0438\u0433\u0438\u043d\u0430\u043b",
        ),
        (
            "document_status",
            "duplicate",
            "\u0414\u0443\u0431\u043b\u0438\u043a\u0430\u0442",
        ),
        (
            "loss_confirmation",
            "No",
            "\u041d\u0435\u0442",
        ),
        (
            "destruction_confirmation",
            "yes",
            "\u0414\u0430",
        ),
        (
            "po_program_type",
            "initial_training",
            (
                "\u041f\u0440\u043e\u0433\u0440\u0430\u043c\u043c\u0430 "
                "\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u043e\u0439 "
                "\u043f\u043e\u0434\u0433\u043e\u0442\u043e\u0432\u043a\u0438 \u043f\u043e "
                "\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u0438 \u0440\u0430\u0431\u043e\u0447\u0435\u0433\u043e, "
                "\u0434\u043e\u043b\u0436\u043d\u043e\u0441\u0442\u0438 \u0441\u043b\u0443\u0436\u0430\u0449\u0435\u0433\u043e"
            ),
        ),
        (
            "sex",
            "male",
            "\u041c\u0443\u0436",
        ),
        (
            "study_form",
            "Full-time",
            "\u041e\u0447\u043d\u0430\u044f",
        ),
        (
            "funding_source",
            "Paid",
            "\u041f\u043b\u0430\u0442\u043d\u043e\u0435 \u043e\u0431\u0443\u0447\u0435\u043d\u0438\u0435",
        ),
        (
            "education_delivery_form",
            "onsite",
            "\u0432 \u043e\u0431\u0440\u0430\u0437\u043e\u0432\u0430\u0442\u0435\u043b\u044c\u043d\u043e\u0439 \u043e\u0440\u0433\u0430\u043d\u0438\u0437\u0430\u0446\u0438\u0438",
        ),
    ],
)
def test_supported_internal_aliases_normalize_to_template(
    classifier: str,
    source: str,
    expected: str,
) -> None:
    assert (
        normalize_frdo_po_classifier_value(
            classifier,
            source,
        )
        == expected
    )


def test_exact_template_classifier_value_is_preserved() -> None:
    value = (
        "\u0420\u0435\u0433\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u044b\u0439 "
        "\u0431\u044e\u0434\u0436\u0435\u0442"
    )

    assert (
        normalize_frdo_po_classifier_value(
            "funding_source",
            value,
        )
        == value
    )


def test_generic_budget_fails_closed() -> None:
    with pytest.raises(
        FrdoPoPortalMappingError,
        match="unsupported",
    ):
        normalize_frdo_po_classifier_value(
            "funding_source",
            "budget",
        )


def test_unknown_profession_fails_closed() -> None:
    with pytest.raises(
        FrdoPoPortalMappingError,
        match="unsupported",
    ):
        normalize_frdo_po_classifier_value(
            "po_profession",
            "Worker",
        )


def test_document_type_uses_frdo_context_classifier() -> None:
    column = FRDO_PO_COLUMNS[0]

    assert column.letter == "A"
    assert column.category == "classifier"
    assert (
        column.source_path
        == "frdo_context.po_document_type"
    )
    assert column.classifier == "document_type"

    assert not any(
        key == "document_type"
        for key, _ in FRDO_PO_KNOWN_MAPPING_GAPS
    )


def test_ai_column_is_constant_blank_contract() -> None:
    column = FRDO_PO_COLUMNS[-1]

    assert column.letter == "AI"
    assert column.category == CATEGORY_CONSTANT
    assert column.source_path is None
    assert "blank" in str(
        column.condition
    ).casefold()


def test_budget_ambiguity_is_explicit_mapping_gap() -> None:
    assert any(
        key == "funding_source_budget"
        for key, _ in FRDO_PO_KNOWN_MAPPING_GAPS
    )


def test_exact_po_document_type_is_supported() -> None:
    value = (
        "\u0421\u0432\u0438\u0434\u0435\u0442\u0435\u043b\u044c\u0441\u0442\u0432\u043e "
        "\u043e \u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u0438 "
        "\u0440\u0430\u0431\u043e\u0447\u0435\u0433\u043e, "
        "\u0434\u043e\u043b\u0436\u043d\u043e\u0441\u0442\u0438 "
        "\u0441\u043b\u0443\u0436\u0430\u0449\u0435\u0433\u043e"
    )

    assert normalize_frdo_po_classifier_value(
        "document_type",
        value,
    ) == value


def test_generic_certificate_document_type_fails_closed() -> None:
    with pytest.raises(
        FrdoPoPortalMappingError,
        match="unsupported",
    ):
        normalize_frdo_po_classifier_value(
            "document_type",
            "Certificate",
        )


def test_art_training_reference_gap_is_explicit() -> None:
    gaps = dict(
        FRDO_PO_KNOWN_MAPPING_GAPS
    )

    assert (
        "training_reference_art_preprofessional_profession"
        in gaps
    )

    message = gaps[
        "training_reference_art_preprofessional_profession"
    ]

    assert "absent" in message
    assert "po_profession" in message
    assert "must not infer or invent" in message


def test_art_training_reference_profession_is_not_in_pinned_classifier() -> None:
    expected = (
        "\u0414\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u044b\u0435 "
        "\u043f\u0440\u0435\u0434\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u044b\u0435 "
        "\u043f\u0440\u043e\u0433\u0440\u0430\u043c\u043c\u044b "
        "\u0432 "
        "\u043e\u0431\u043b\u0430\u0441\u0442\u0438 "
        "\u0438\u0441\u043a\u0443\u0441\u0441\u0442\u0432"
    )

    values = get_frdo_po_classifier_values(
        "po_profession"
    )

    assert expected not in values


def test_art_training_reference_profession_is_not_guessed() -> None:
    expected = (
        "\u0414\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u044b\u0435 "
        "\u043f\u0440\u0435\u0434\u043f\u0440\u043e\u0444\u0435\u0441\u0441\u0438\u043e\u043d\u0430\u043b\u044c\u043d\u044b\u0435 "
        "\u043f\u0440\u043e\u0433\u0440\u0430\u043c\u043c\u044b "
        "\u0432 "
        "\u043e\u0431\u043b\u0430\u0441\u0442\u0438 "
        "\u0438\u0441\u043a\u0443\u0441\u0441\u0442\u0432"
    )

    with pytest.raises(
        FrdoPoPortalMappingError
    ):
        normalize_frdo_po_classifier_value(
            "po_profession",
            expected,
        )
