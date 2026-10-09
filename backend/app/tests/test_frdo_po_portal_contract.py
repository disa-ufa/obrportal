from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

import app.services.frdo_po_portal_contract as contract_module
from app.services.compliance_registry_contract import (
    PROGRAM_TYPE_DPO_ADVANCED_TRAINING,
    PROGRAM_TYPE_DPO_PROFESSIONAL_RETRAINING,
    PROGRAM_TYPE_VOCATIONAL_TRAINING,
)
from app.services.frdo_po_portal_contract import (
    FRDO_PO_PORTAL_CONTRACT_STATUS_WORKING_REFERENCE,
    FRDO_PO_TEMPLATE_CONTRACT_VERSION,
    FRDO_PO_TEMPLATE_KIND,
    FRDO_PO_TEMPLATE_SHA256,
    FRDO_PO_TEMPLATE_SOURCE_REFERENCE,
    FRDO_PO_XLSX_EXTENSION,
    FRDO_PO_XLSX_FORMAT,
    FRDO_PO_XLSX_MIME_TYPE,
    FrdoPoPortalContractError,
    FrdoPoPortalContractUnavailable,
    get_frdo_po_portal_contract,
    read_verified_frdo_po_template_bytes,
    require_frdo_po_portal_contract,
)


def test_frdo_po_contract_is_working_reference() -> None:
    contract = get_frdo_po_portal_contract(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    assert (
        contract.program_type
        == PROGRAM_TYPE_VOCATIONAL_TRAINING
    )
    assert contract.template_kind == FRDO_PO_TEMPLATE_KIND
    assert contract.template_kind == "po"
    assert (
        contract.status
        == FRDO_PO_PORTAL_CONTRACT_STATUS_WORKING_REFERENCE
    )
    assert contract.is_working_reference is True
    assert contract.is_available is True
    assert contract.is_confirmed is False


def test_frdo_po_contract_has_pinned_working_reference_metadata() -> None:
    contract = get_frdo_po_portal_contract(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    assert (
        contract.template_contract_version
        == FRDO_PO_TEMPLATE_CONTRACT_VERSION
    )
    assert (
        contract.template_contract_version
        == "frdo-po-working-reference-v1"
    )
    assert (
        contract.source_reference
        == FRDO_PO_TEMPLATE_SOURCE_REFERENCE
    )
    assert (
        contract.source_reference
        == "frdo_po_working_reference_v1.xlsx"
    )
    assert contract.source_sha256 == FRDO_PO_TEMPLATE_SHA256
    assert (
        contract.source_sha256
        == "205853ef3cfcf58626d6cdd9cf070d1b1b178792f09d1b222e5064c16371f82a"
    )
    assert contract.file_format == FRDO_PO_XLSX_FORMAT
    assert contract.file_format == "xlsx"
    assert contract.mime_type == FRDO_PO_XLSX_MIME_TYPE
    assert (
        contract.mime_type
        == "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    )
    assert contract.extension == FRDO_PO_XLSX_EXTENSION
    assert contract.extension == ".xlsx"


def test_frdo_po_contract_requirement_accepts_working_reference() -> None:
    contract = require_frdo_po_portal_contract(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    assert contract.is_working_reference is True
    assert contract.is_confirmed is False


def test_frdo_po_template_resource_matches_pinned_sha() -> None:
    content = read_verified_frdo_po_template_bytes(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    assert (
        hashlib.sha256(content).hexdigest()
        == FRDO_PO_TEMPLATE_SHA256
    )


def test_frdo_po_template_basic_workbook_contract() -> None:
    content = read_verified_frdo_po_template_bytes(
        program_type=PROGRAM_TYPE_VOCATIONAL_TRAINING,
    )

    workbook = load_workbook(
        BytesIO(content),
        read_only=False,
        data_only=False,
    )

    assert workbook.sheetnames == [
        "Шаблон",
        "Проверки",
    ]

    template_sheet = workbook[
        "Шаблон"
    ]
    checks_sheet = workbook[
        "Проверки"
    ]

    assert template_sheet.max_column == 35
    assert template_sheet.max_row == 1002
    assert checks_sheet.sheet_state == "hidden"
    assert checks_sheet.max_column == 19
    assert checks_sheet.max_row == 5600
    assert (
        len(
            template_sheet
            .data_validations
            .dataValidation
        )
        == 20
    )


def test_frdo_po_template_checksum_mismatch_fails_closed(
    tmp_path,
    monkeypatch,
) -> None:
    changed_template = (
        tmp_path
        / FRDO_PO_TEMPLATE_SOURCE_REFERENCE
    )
    changed_template.write_bytes(b"changed")

    monkeypatch.setattr(
        contract_module,
        "FRDO_PO_TEMPLATE_RESOURCE_PATH",
        changed_template,
    )

    with pytest.raises(
        FrdoPoPortalContractUnavailable,
        match="template checksum mismatch",
    ):
        read_verified_frdo_po_template_bytes(
            program_type=(
                PROGRAM_TYPE_VOCATIONAL_TRAINING
            ),
        )


@pytest.mark.parametrize(
    "program_type",
    [
        PROGRAM_TYPE_DPO_ADVANCED_TRAINING,
        PROGRAM_TYPE_DPO_PROFESSIONAL_RETRAINING,
    ],
)
def test_frdo_po_contract_rejects_dpo_programs(
    program_type: str,
) -> None:
    with pytest.raises(
        FrdoPoPortalContractError,
        match="supports only vocational_training",
    ):
        get_frdo_po_portal_contract(
            program_type=program_type,
        )

    with pytest.raises(
        FrdoPoPortalContractError,
        match="supports only vocational_training",
    ):
        read_verified_frdo_po_template_bytes(
            program_type=program_type,
        )
